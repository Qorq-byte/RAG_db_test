"""PowerPoint parser preserving slide numbers and titles."""

from pptx import Presentation

from ragdb.domain.enums import SourceType
from ragdb.domain.errors import DocumentParseError
from ragdb.domain.models import Document, DocumentUnit, Source, SourcePosition
from ragdb.infrastructure.parsers._common import source_path


class PowerPointParser:
    name = "powerpoint"
    version = "1.0"

    def supports(self, source: Source) -> bool:
        return source.source_type is SourceType.POWERPOINT

    def parse(self, source: Source) -> Document:
        path = source_path(source)
        try:
            presentation = Presentation(path)
        except Exception as exc:
            raise DocumentParseError(source.uri, str(exc)) from exc

        units: list[DocumentUnit] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        parts.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip(" |"):
                            parts.append(row_text)
            if not parts:
                continue
            title_shape = slide.shapes.title
            slide_title = title_shape.text.strip() if title_shape is not None else ""
            metadata: dict[str, str | int] = {"shape_text_count": len(parts)}
            if slide_title:
                metadata["slide_title"] = slide_title
            units.append(
                DocumentUnit(
                    text="\n".join(parts),
                    position=SourcePosition(slide=slide_number),
                    metadata=metadata,
                )
            )

        if not units:
            raise DocumentParseError(source.uri, "演示文稿不包含可解析文本")
        native_title = (presentation.core_properties.title or "").strip()
        metadata: dict[str, str | int] = {
            "format": "pptx",
            "filename": path.name,
            "slide_count": len(presentation.slides),
        }
        if native_title:
            metadata["native_title"] = native_title
        return Document(source_id=source.id, title=source.title, units=tuple(units), metadata=metadata)

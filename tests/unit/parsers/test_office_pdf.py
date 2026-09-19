import pymupdf
import pytest
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches

from ragdb.domain.enums import SourceType
from ragdb.domain.errors import OcrRequiredError
from ragdb.infrastructure.parsers.pdf import PdfParser
from ragdb.infrastructure.parsers.powerpoint import PowerPointParser
from ragdb.infrastructure.parsers.word import WordParser


def test_pdf_parser_preserves_pages_and_flags_sparse_pages(tmp_path, make_source) -> None:
    path = tmp_path / "notes.pdf"
    with pymupdf.open() as pdf:
        first = pdf.new_page()
        first.insert_text((72, 72), "A sufficiently long searchable PDF text layer for parser testing.")
        pdf.new_page()
        pdf.save(path)

    document = PdfParser(min_characters_per_page=20).parse(
        make_source(path, SourceType.PDF)
    )
    assert document.units[0].position.page == 1
    assert document.metadata["ocr_suggested"] is True
    assert document.metadata["ocr_pages"] == [2]


def test_pdf_parser_requests_ocr_when_all_pages_are_sparse(tmp_path, make_source) -> None:
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(path)

    with pytest.raises(OcrRequiredError) as error:
        PdfParser().parse(make_source(path, SourceType.PDF))
    assert error.value.page_numbers == (1,)


def test_word_parser_preserves_headings_and_tables(tmp_path, make_source) -> None:
    path = tmp_path / "notes.docx"
    word = WordDocument()
    word.core_properties.title = "原生标题"
    word.add_heading("第一章", level=1)
    word.add_paragraph("正文内容")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "键"
    table.cell(0, 1).text = "值"
    word.save(path)

    document = WordParser().parse(make_source(path, SourceType.WORD))
    assert document.title == "测试资料"
    assert document.metadata["native_title"] == "原生标题"
    assert document.units[1].position.heading_path == ("第一章",)
    assert document.units[2].metadata["block_type"] == "table"


def test_powerpoint_parser_preserves_slide_number_and_title(tmp_path, make_source) -> None:
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "架构"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    textbox.text = "解析、分块、索引"
    deck.save(path)

    document = PowerPointParser().parse(make_source(path, SourceType.POWERPOINT))
    assert document.units[0].position.slide == 1
    assert document.units[0].metadata["slide_title"] == "架构"
    assert "解析、分块、索引" in document.units[0].text
